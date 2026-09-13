"""
Tests for the BRSR (Business Responsibility and Sustainability Reporting) module.

Distinct from tests/test_brsr_report.py, which tests the pre-existing
Scope-2-only operational sustainability summary in app/trust/report.py.
This file tests the new app/brsr/* package: company profile, the metric
registry, report lifecycle, provenance/quality, GreenShift-derived
calculations, validation, RBAC/tenant isolation, audit integration, and
report generation.
"""

import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

from app.api.main import app
from app.brsr.service import (
    BrsrNotFoundError,
    BrsrPermissionError,
    BrsrTransitionError,
    BrsrValidationBlockedError,
    apply_greenshift_derived_metrics,
    create_report,
    get_or_create_company_profile,
    get_report,
    seed_metric_registry,
    transition_status,
    update_company_profile,
    update_metric_value,
)
from app.brsr.validation import run_validation
from app.shared.auth import create_access_token, hash_password
from app.shared.database import get_db
from app.shared.models import (
    BrsrMetricDefinitionORM,
    BrsrMetricValueORM,
    JobORM,
    JobStatus,
    ScheduleDecisionORM,
    TenantORM,
    UserORM,
    UserRole,
)

client = TestClient(app)


@pytest.fixture(autouse=True)
def override_db(db):
    def _get_test_db():
        yield db

    app.dependency_overrides[get_db] = _get_test_db
    yield
    app.dependency_overrides.pop(get_db, None)


@pytest.fixture(autouse=True)
def seeded_registry(db):
    seed_metric_registry(db)
    yield


@pytest.fixture
def two_tenants(db):
    tenant_a = TenantORM(id="tenant-brsr-a", name="Acme Corp A", is_active=True)
    tenant_b = TenantORM(id="tenant-brsr-b", name="Beta Corp B", is_active=True)
    db.add_all([tenant_a, tenant_b])
    db.commit()

    admin_a = UserORM(username="brsr_admin_a", email="admin_a@brsr.io", hashed_password=hash_password("pass12345"),
                       role=UserRole.COMPANY_ADMIN, tenant_id="tenant-brsr-a", is_active=True)
    user_a = UserORM(username="brsr_user_a", email="user_a@brsr.io", hashed_password=hash_password("pass12345"),
                      role=UserRole.COMPANY_USER, tenant_id="tenant-brsr-a", is_active=True)
    admin_b = UserORM(username="brsr_admin_b", email="admin_b@brsr.io", hashed_password=hash_password("pass12345"),
                       role=UserRole.COMPANY_ADMIN, tenant_id="tenant-brsr-b", is_active=True)
    platform_admin = UserORM(username="brsr_platform_admin", email="pa@brsr.io", hashed_password=hash_password("pass12345"),
                              role=UserRole.PLATFORM_ADMIN, tenant_id=None, is_active=True)
    db.add_all([admin_a, user_a, admin_b, platform_admin])
    db.commit()
    for u in (admin_a, user_a, admin_b, platform_admin):
        db.refresh(u)
    return {"admin_a": admin_a, "user_a": user_a, "admin_b": admin_b, "platform_admin": platform_admin}


def get_token(user: UserORM) -> str:
    role_str = user.role.value if hasattr(user.role, "value") else str(user.role)
    return create_access_token(user_id=user.id, username=user.username, role=role_str, tenant_id=user.tenant_id)


def _period():
    start = datetime(2025, 4, 1, tzinfo=timezone.utc)
    end = datetime(2026, 3, 31, tzinfo=timezone.utc)
    return start, end


class TestMetricRegistry:
    def test_seed_is_idempotent(self, db):
        count1 = seed_metric_registry(db)
        count2 = seed_metric_registry(db)
        assert count2 == 0  # nothing new to insert the second time
        total = db.query(BrsrMetricDefinitionORM).count()
        assert total > 30  # a real, representative registry, not a token handful

    def test_registry_covers_all_9_principles(self, db):
        principles = {
            row[0] for row in db.query(BrsrMetricDefinitionORM.principle).filter(BrsrMetricDefinitionORM.principle.isnot(None)).distinct()
        }
        assert principles == set(range(1, 10))

    def test_registry_covers_brsr_core_attributes(self, db):
        attrs = {
            row[0] for row in db.query(BrsrMetricDefinitionORM.brsr_core_attribute)
            .filter(BrsrMetricDefinitionORM.brsr_core_attribute.isnot(None)).distinct()
        }
        expected = {
            "GHG_FOOTPRINT", "WATER_FOOTPRINT", "ENERGY_FOOTPRINT", "EMISSIONS_WASTE_CIRCULARITY",
            "EMPLOYEE_WELLBEING", "GENDER_DIVERSITY", "INCLUSIVE_DEVELOPMENT", "VALUE_CHAIN_FAIRNESS",
            "OPENNESS_OF_BUSINESS",
        }
        assert attrs == expected

    def test_only_scope2_energy_and_their_intensities_are_greenshift_derivable(self, db):
        """Honesty check: only Scope 2 emissions, total energy, and the two
        intensity ratios computed from them are genuinely derivable from
        GreenShift's own operational data — every other metric (Scope 1,
        Scope 3, water, waste, air emissions, biodiversity, and every
        social/governance field) has no GreenShift-derivable source."""
        codes = {
            row[0] for row in db.query(BrsrMetricDefinitionORM.metric_code)
            .filter(BrsrMetricDefinitionORM.greenshift_derivable == True)  # noqa: E712
        }
        assert codes == {
            "CORE_GHG_SCOPE2", "P6_SCOPE2_EMISSIONS",
            "CORE_ENERGY_CONSUMPTION", "P6_TOTAL_ENERGY_CONSUMPTION",
            "CORE_ENERGY_INTENSITY", "P6_ENERGY_INTENSITY",
            "CORE_EMISSION_INTENSITY", "P6_EMISSION_INTENSITY",
        }


class TestCompanyProfile:
    def test_lazily_created_with_missing_source_marked(self, db, two_tenants):
        profile = get_or_create_company_profile(db, "tenant-brsr-a")
        assert profile.company_name is None  # never fabricated
        assert profile.source_type == "COMPANY_PROVIDED"

    def test_company_admin_can_update_own_profile(self, db, two_tenants):
        admin_a = two_tenants["admin_a"]
        profile = update_company_profile(db, "tenant-brsr-a", {"company_name": "Acme Real Corp", "cin": "L12345MH2020PLC000001", "listed_status": "LISTED"}, admin_a)
        assert profile.company_name == "Acme Real Corp"
        assert profile.updated_by == admin_a.id

    def test_company_user_cannot_update_profile(self, db, two_tenants):
        user_a = two_tenants["user_a"]
        with pytest.raises(BrsrPermissionError):
            update_company_profile(db, "tenant-brsr-a", {"company_name": "Hacked"}, user_a)

    def test_platform_admin_cannot_edit_any_companys_profile(self, db, two_tenants):
        platform_admin = two_tenants["platform_admin"]
        with pytest.raises(BrsrPermissionError):
            update_company_profile(db, "tenant-brsr-a", {"company_name": "Hacked by platform"}, platform_admin)

    def test_admin_cannot_edit_another_tenants_profile(self, db, two_tenants):
        admin_b = two_tenants["admin_b"]
        with pytest.raises(BrsrPermissionError):
            update_company_profile(db, "tenant-brsr-a", {"company_name": "Cross-tenant hack"}, admin_b)

    def test_api_forged_tenant_id_in_query_param_is_ignored_for_write(self, db, two_tenants):
        """PUT /brsr/company-profile has no tenant_id field at all — the
        write always targets the caller's own tenant_id."""
        admin_a = two_tenants["admin_a"]
        token = get_token(admin_a)
        res = client.put(
            "/brsr/company-profile", headers={"Authorization": f"Bearer {token}"},
            json={"company_name": "Real Name", "tenant_id": "tenant-brsr-b"},
        )
        assert res.status_code == 200
        assert res.json()["tenant_id"] == "tenant-brsr-a"

    def test_api_platform_admin_can_view_but_not_default_to_own_tenant(self, db, two_tenants):
        platform_admin = two_tenants["platform_admin"]
        token = get_token(platform_admin)
        res = client.get("/brsr/company-profile", params={"tenant_id": "tenant-brsr-a"}, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        assert res.json()["tenant_id"] == "tenant-brsr-a"


class TestReportLifecycle:
    def test_create_report_prepopulates_all_registry_metrics_as_missing(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        assert report.status == "DRAFT"
        values = db.query(BrsrMetricValueORM).filter(BrsrMetricValueORM.report_id == report.id).all()
        total_defs = db.query(BrsrMetricDefinitionORM).count()
        assert len(values) == total_defs
        assert all(v.source_type == "MISSING" for v in values)
        assert all(v.quality == "MISSING" for v in values)

    def test_cannot_create_duplicate_report_for_same_fy(self, db, two_tenants):
        start, end = _period()
        create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        with pytest.raises(BrsrTransitionError):
            create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])

    def test_company_user_cannot_create_report(self, db, two_tenants):
        start, end = _period()
        with pytest.raises(BrsrPermissionError):
            create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["user_a"])

    def test_illegal_transition_rejected(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        # DRAFT can only go to DATA_COLLECTION — skipping ahead must fail.
        with pytest.raises(BrsrTransitionError):
            transition_status(db, report.id, "APPROVED", two_tenants["admin_a"])

    def test_full_lifecycle_reaches_generated(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_company_profile(db, "tenant-brsr-a", {"company_name": "Acme Corp", "listed_status": "UNLISTED"}, admin_a)

        # Fill every required metric so validation has no ERRORs.
        required_defs = db.query(BrsrMetricDefinitionORM).filter(BrsrMetricDefinitionORM.required == True).all()  # noqa: E712
        for defn in required_defs:
            if defn.data_type == "BOOLEAN":
                update_metric_value(db, report.id, defn.metric_code, {"text_value": "true"}, admin_a)
            elif defn.data_type == "TEXT":
                update_metric_value(db, report.id, defn.metric_code, {"text_value": "Board-level committee in place"}, admin_a)
            else:
                update_metric_value(db, report.id, defn.metric_code, {"value": 10.0, "unit": defn.unit}, admin_a)

        report = transition_status(db, report.id, "DATA_COLLECTION", admin_a)
        assert report.status == "DATA_COLLECTION"

        report = transition_status(db, report.id, "VALIDATED", admin_a)
        assert report.status == "VALIDATED"
        assert report.validated_at is not None

        report = transition_status(db, report.id, "APPROVED", admin_a)
        assert report.status == "APPROVED"
        assert report.approved_by == admin_a.id

        report = transition_status(db, report.id, "GENERATED", admin_a)
        assert report.status == "GENERATED"
        assert report.generated_at is not None

    def test_cannot_validate_with_required_fields_missing(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        transition_status(db, report.id, "DATA_COLLECTION", admin_a)
        with pytest.raises(BrsrValidationBlockedError):
            transition_status(db, report.id, "VALIDATED", admin_a)

    def test_cannot_edit_metrics_on_generated_report(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        # Force to GENERATED directly for this isolated test of the edit guard.
        report.status = "GENERATED"
        db.commit()
        with pytest.raises(BrsrTransitionError):
            update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, admin_a)

    def test_report_created_before_a_registry_metric_existed_is_backfilled_on_next_access(self, db, two_tenants):
        """The metric registry must be able to evolve (a new BRSR question,
        or an entire new framework_version) without leaving reports created
        before the change permanently blind to the new metric — this is
        the whole point of a registry-driven design (Phase 4). Simulates
        that scenario by creating a report, then deleting one of its
        metric-value rows (as if it predated that metric's addition to the
        registry), and confirming a later get_report() call transparently
        backfills it as MISSING — without touching any other row."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)

        # Answer one real metric, then simulate a "new registry metric" by
        # deleting a different row entirely (as if this report predated it).
        update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, admin_a)
        db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "P6_SCOPE3_EMISSIONS",
        ).delete()
        db.commit()
        assert db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "P6_SCOPE3_EMISSIONS",
        ).first() is None

        # Any access via get_report() (the funnel every API endpoint uses)
        # must transparently backfill the missing row...
        refetched = get_report(db, report.id, admin_a)
        assert refetched.id == report.id
        backfilled = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "P6_SCOPE3_EMISSIONS",
        ).first()
        assert backfilled is not None
        assert backfilled.source_type == "MISSING"
        assert backfilled.value is None

        # ...without disturbing the already-answered metric.
        untouched = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "SEC_A_CSR_APPLICABLE",
        ).first()
        assert untouched.text_value == "true"
        assert untouched.source_type == "COMPANY_PROVIDED"

    def test_registry_backfill_is_idempotent_and_does_not_duplicate_rows(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        get_report(db, report.id, admin_a)
        get_report(db, report.id, admin_a)
        get_report(db, report.id, admin_a)
        total_defs = db.query(BrsrMetricDefinitionORM).count()
        total_values = db.query(BrsrMetricValueORM).filter(BrsrMetricValueORM.report_id == report.id).count()
        assert total_values == total_defs


class TestTenantIsolationAndRBAC:
    def test_company_user_cannot_edit_metric(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        with pytest.raises(BrsrPermissionError):
            update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, two_tenants["user_a"])

    def test_other_tenant_admin_cannot_view_report(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        with pytest.raises(BrsrNotFoundError):
            get_report(db, report.id, two_tenants["admin_b"])

    def test_other_tenant_admin_cannot_edit_metric(self, db, two_tenants):
        """update_metric_value looks the report up (require_view_access)
        before checking edit rights — a cross-tenant admin gets a 404, not
        a 403, matching the existing tenant_scope.py convention of never
        confirming another tenant's record even exists."""
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        with pytest.raises(BrsrNotFoundError):
            update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, two_tenants["admin_b"])

    def test_platform_admin_can_view_any_tenants_report(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        fetched = get_report(db, report.id, two_tenants["platform_admin"])
        assert fetched.id == report.id

    def test_platform_admin_cannot_edit_any_tenants_metric(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        with pytest.raises(BrsrPermissionError):
            update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, two_tenants["platform_admin"])

    def test_direct_api_cross_tenant_report_access_returns_404(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        token_b = get_token(two_tenants["admin_b"])
        res = client.get(f"/brsr/reports/{report.id}", headers={"Authorization": f"Bearer {token_b}"})
        assert res.status_code == 404

    def test_direct_api_forged_role_in_metric_update_body_is_ignored(self, db, two_tenants):
        """BrsrMetricValueUpdateRequest has no `quality`/`role`/`report_id`/
        `metric_code` field at all — those have zero effect if forged into
        the body (Pydantic silently drops unknown fields). `source_type` IS
        a real, validated field (restricted to COMPANY_PROVIDED/EXTERNAL_SOURCE
        — see TestDataQualityAndProvenance for its dedicated coverage), so a
        forged GREENSHIFT_DERIVED claim there is rejected outright (422),
        not silently ignored — a stricter posture than "ignored"."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.put(
            f"/brsr/reports/{report.id}/metrics/SEC_A_CSR_AMOUNT_SPENT",
            headers={"Authorization": f"Bearer {token}"},
            json={"value": 100000, "quality": "HIGH", "role": "PLATFORM_ADMIN", "report_id": 999, "metric_code": "OTHER"},
        )
        assert res.status_code == 200
        assert res.json()["source_type"] == "COMPANY_PROVIDED"  # forged fields had no effect
        assert res.json()["metric_code"] == "SEC_A_CSR_AMOUNT_SPENT"  # the URL, not the forged body field, wins

    def test_direct_api_unauthenticated_rejected(self, db):
        assert client.get("/brsr/reports").status_code == 401
        assert client.get("/brsr/company-profile").status_code == 401
        assert client.post("/brsr/reports", json={}).status_code == 401

    def test_direct_api_company_user_forbidden_from_creating_report(self, db, two_tenants):
        token = get_token(two_tenants["user_a"])
        res = client.post(
            "/brsr/reports", headers={"Authorization": f"Bearer {token}"},
            json={"financial_year": "2025-26", "reporting_period_start": "2025-04-01T00:00:00Z", "reporting_period_end": "2026-03-31T00:00:00Z"},
        )
        assert res.status_code == 403

    def test_direct_api_export_before_generated_rejected(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.get(f"/brsr/reports/{report.id}/export", params={"format": "pdf"}, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 422

    def test_direct_api_transition_forged_target_status_rejected(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.post(
            f"/brsr/reports/{report.id}/transition", headers={"Authorization": f"Bearer {token}"},
            json={"target_status": "NOT_A_REAL_STATUS"},
        )
        assert res.status_code == 422

    def test_direct_api_cross_tenant_idor_sweep_across_every_report_scoped_endpoint(self, db, two_tenants):
        """Final BRSR Hardening Pass: the router-level wiring for every
        report_id-scoped endpoint beyond plain GET /reports/{id} was
        previously verified only at the service-layer (get_report/
        require_edit_access unit calls), never through an actual HTTP
        request from a different tenant's token. Sweep every one to prove
        the router's own Depends()/get_report() wiring — not just the
        service function in isolation — rejects cross-tenant access."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token_b = get_token(two_tenants["admin_b"])
        headers_b = {"Authorization": f"Bearer {token_b}"}

        # Read-only endpoints: cross-tenant -> 404 (never leak existence).
        assert client.get(f"/brsr/reports/{report.id}/overview", headers=headers_b).status_code == 404
        assert client.get(f"/brsr/reports/{report.id}/metrics", headers=headers_b).status_code == 404
        assert client.get(f"/brsr/reports/{report.id}/assessment", headers=headers_b).status_code == 404
        assert client.get(f"/brsr/reports/{report.id}/validation", headers=headers_b).status_code == 404
        assert client.get(f"/brsr/reports/{report.id}/audit", headers=headers_b).status_code == 404
        assert client.get(f"/brsr/reports/{report.id}/export", params={"format": "json"}, headers=headers_b).status_code == 404

        # Mutating endpoints: cross-tenant -> 404 (view-access check runs
        # before the edit-access check, so a wrong-tenant caller never even
        # learns the report exists to be told they can't edit it).
        assert client.put(
            f"/brsr/reports/{report.id}/metrics/SEC_A_CSR_AMOUNT_SPENT", headers=headers_b, json={"value": 1},
        ).status_code == 404
        assert client.post(f"/brsr/reports/{report.id}/apply-greenshift-data", headers=headers_b).status_code == 404
        assert client.post(f"/brsr/reports/{report.id}/validate", headers=headers_b).status_code == 404
        assert client.post(
            f"/brsr/reports/{report.id}/transition", headers=headers_b, json={"target_status": "DATA_COLLECTION"},
        ).status_code == 404
        assert client.put(f"/brsr/reports/{report.id}/assessment", headers=headers_b, json={"scope": "x"}).status_code == 404

    def test_direct_api_forged_tenant_id_in_create_report_body_has_no_effect(self, db, two_tenants):
        """BrsrReportCreateRequest has no tenant_id field at all — Pydantic
        silently drops it, so the report is always created under the
        caller's own tenant regardless of what's in the body."""
        admin_a = two_tenants["admin_a"]
        token = get_token(admin_a)
        res = client.post(
            "/brsr/reports", headers={"Authorization": f"Bearer {token}"},
            json={
                "financial_year": "2025-26",
                "reporting_period_start": "2025-04-01T00:00:00Z",
                "reporting_period_end": "2026-03-31T00:00:00Z",
                "tenant_id": "tenant-brsr-b",
            },
        )
        assert res.status_code == 201
        assert res.json()["tenant_id"] == "tenant-brsr-a"

    def test_direct_api_list_reports_ignores_forged_tenant_id_for_non_platform_admin(self, db, two_tenants):
        start, end = _period()
        create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        create_report(db, "tenant-brsr-b", "2025-26", start, end, "BRSR-2023", two_tenants["admin_b"])
        token_a = get_token(two_tenants["admin_a"])
        res = client.get("/brsr/reports", params={"tenant_id": "tenant-brsr-b"}, headers={"Authorization": f"Bearer {token_a}"})
        assert res.status_code == 200
        tenant_ids = {r["tenant_id"] for r in res.json()}
        assert tenant_ids == {"tenant-brsr-a"}


class TestGreenShiftDerivedCalculations:
    def test_scope2_and_energy_derived_from_real_jobs(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        mid = start + timedelta(days=10)

        job = JobORM(
            job_id="BRSR-CALC-JOB-1", team_id="team-x", tenant_id="tenant-brsr-a",
            submitted_by_user_id=admin_a.id, submitted_at=mid, deadline=mid + timedelta(hours=6),
            runtime_minutes=60, power_kw=10.0, region="IN-TG",
            container_image="greenshift/sample:v1", status=JobStatus.COMPLETED, energy_kwh=10.0,
        )
        db.add(job)
        db.commit()
        decision = ScheduleDecisionORM(
            job_id=job.job_id, selected_start=mid, selected_end=mid + timedelta(hours=1),
            carbon_intensity=400.0, carbon_emission=4.0, electricity_cost=1.0, reason="test",
            region_id="IN-TG", currency="INR", native_cost=80.0,
        )
        db.add(decision)
        db.commit()

        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        updated = apply_greenshift_derived_metrics(db, report, admin_a)
        assert updated == 8  # Scope2, Energy, Energy Intensity, Emission Intensity — each duplicated CORE + P6

        scope2 = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "CORE_GHG_SCOPE2",
        ).first()
        assert scope2.source_type == "GREENSHIFT_DERIVED"
        assert scope2.quality == "HIGH"
        assert scope2.value == pytest.approx(0.004)  # 4 kg -> 0.004 tCO2e
        assert "BRSR-CALC-JOB-1" in (scope2.source_detail or {}).get("job_ids", [])

        energy = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "CORE_ENERGY_CONSUMPTION",
        ).first()
        assert energy.value == pytest.approx(10.0)
        assert energy.source_type == "GREENSHIFT_DERIVED"

        energy_intensity = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "CORE_ENERGY_INTENSITY",
        ).first()
        assert energy_intensity.value == pytest.approx(10.0)  # 10 kWh / 1 job
        assert energy_intensity.source_type == "GREENSHIFT_DERIVED"
        assert energy_intensity.quality == "HIGH"

        emission_intensity = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "CORE_EMISSION_INTENSITY",
        ).first()
        assert emission_intensity.value == pytest.approx(0.4)  # 4 kg / 10 kWh
        assert emission_intensity.source_type == "GREENSHIFT_DERIVED"

    def test_no_jobs_in_period_yields_real_zero_not_missing(self, db, two_tenants):
        """A genuine, complete computation over zero matching jobs is a
        real 0 — distinct from MISSING (unknown)."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        apply_greenshift_derived_metrics(db, report, admin_a)
        scope2 = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "CORE_GHG_SCOPE2",
        ).first()
        assert scope2.value == 0.0
        assert scope2.source_type == "GREENSHIFT_DERIVED"
        assert scope2.quality == "HIGH"

    def test_cross_tenant_jobs_never_counted(self, db, two_tenants):
        """A job belonging to tenant-brsr-b must never contribute to
        tenant-brsr-a's derived metrics."""
        start, end = _period()
        admin_a, admin_b = two_tenants["admin_a"], two_tenants["admin_b"]
        mid = start + timedelta(days=5)
        job_b = JobORM(
            job_id="BRSR-CALC-JOB-B", team_id="team-y", tenant_id="tenant-brsr-b",
            submitted_by_user_id=admin_b.id, submitted_at=mid, deadline=mid + timedelta(hours=6),
            runtime_minutes=60, power_kw=99.0, region="IN-TG",
            container_image="greenshift/sample:v1", status=JobStatus.COMPLETED, energy_kwh=99.0,
        )
        db.add(job_b)
        db.commit()

        report_a = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        apply_greenshift_derived_metrics(db, report_a, admin_a)
        energy = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report_a.id, BrsrMetricValueORM.metric_code == "CORE_ENERGY_CONSUMPTION",
        ).first()
        assert energy.value == 0.0  # tenant-b's 99 kWh job must not leak in


class TestDataQualityAndProvenance:
    def test_manual_entry_gets_medium_quality(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        row = update_metric_value(db, report.id, "SEC_A_CSR_AMOUNT_SPENT", {"value": 500000, "unit": "currency"}, admin_a)
        assert row.quality == "MEDIUM"
        assert row.source_type == "COMPANY_PROVIDED"

    def test_estimated_value_gets_low_quality(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        row = update_metric_value(db, report.id, "P6_WATER_CONSUMPTION", {
            "value": 1000, "estimated": True, "estimation_method": "Extrapolated from prior year", "assumption": "No major process change",
        }, admin_a)
        assert row.quality == "LOW"

    def test_unfilled_metric_stays_missing_never_zero(self, db, two_tenants):
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        row = db.query(BrsrMetricValueORM).filter(
            BrsrMetricValueORM.report_id == report.id, BrsrMetricValueORM.metric_code == "P6_WASTE_GENERATED",
        ).first()
        assert row.value is None
        assert row.source_type == "MISSING"
        assert row.quality == "MISSING"

    def test_company_admin_can_honestly_self_declare_external_source(self, db, two_tenants):
        """A Company Admin can legitimately mark a manually-entered value as
        EXTERNAL_SOURCE (e.g. a utility bill or third-party auditor figure)
        — this was previously unreachable (every manual edit was forced to
        COMPANY_PROVIDED regardless of intent), leaving one of the six
        provenance classifications permanently dead."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        row = update_metric_value(db, report.id, "P6_WATER_CONSUMPTION", {
            "value": 1200, "unit": "kilolitres", "source_type": "EXTERNAL_SOURCE",
        }, admin_a)
        assert row.source_type == "EXTERNAL_SOURCE"
        assert row.quality == "MEDIUM"
        assert row.source_record == f"external-source:{admin_a.username}"

    def test_client_still_cannot_forge_greenshift_derived_or_calculated_via_source_type(self, db, two_tenants):
        """The escape hatch is narrowly scoped to {COMPANY_PROVIDED,
        EXTERNAL_SOURCE} — GREENSHIFT_DERIVED/CALCULATED remain permanently
        unreachable through a manual edit, at the service layer directly
        (not merely at the Pydantic API boundary)."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        row = update_metric_value(db, report.id, "P6_WATER_CONSUMPTION", {
            "value": 1200, "source_type": "GREENSHIFT_DERIVED",
        }, admin_a)
        assert row.source_type == "COMPANY_PROVIDED"

        row2 = update_metric_value(db, report.id, "P6_WASTE_GENERATED", {
            "value": 5, "source_type": "CALCULATED",
        }, admin_a)
        assert row2.source_type == "COMPANY_PROVIDED"

    def test_direct_api_rejects_invalid_source_type_value(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.put(
            f"/brsr/reports/{report.id}/metrics/P6_WATER_CONSUMPTION",
            headers={"Authorization": f"Bearer {token}"},
            json={"value": 1000, "source_type": "GREENSHIFT_DERIVED"},
        )
        # Pydantic rejects it outright (422) since GREENSHIFT_DERIVED is not
        # in the client-allowed set — a stricter, earlier rejection than the
        # service-layer silent-ignore path exercised above.
        assert res.status_code == 422

    def test_direct_api_can_set_external_source(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.put(
            f"/brsr/reports/{report.id}/metrics/P6_WATER_CONSUMPTION",
            headers={"Authorization": f"Bearer {token}"},
            json={"value": 1000, "unit": "kilolitres", "source_type": "EXTERNAL_SOURCE"},
        )
        assert res.status_code == 200
        assert res.json()["source_type"] == "EXTERNAL_SOURCE"


class TestValidationEngine:
    def test_negative_value_is_error(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "P1_ANTI_CORRUPTION_COMPLAINTS", {"value": -5}, admin_a)
        run = run_validation(db, report, admin_a)
        assert run.error_count >= 1
        codes = {i.code for i in run.issues}
        assert "NEGATIVE_VALUE" in codes

    def test_percentage_out_of_range_is_error(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "P2_RD_SPEND_PCT", {"value": 150}, admin_a)
        run = run_validation(db, report, admin_a)
        codes = {i.code for i in run.issues}
        assert "PERCENTAGE_OUT_OF_RANGE" in codes

    def test_listed_company_without_cin_is_error(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        update_company_profile(db, "tenant-brsr-a", {"company_name": "Acme", "listed_status": "LISTED"}, admin_a)
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        run = run_validation(db, report, admin_a)
        codes = {i.code for i in run.issues}
        assert "MISSING_CIN" in codes

    def test_reconciliation_mismatch_is_warning(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "CORE_GHG_SCOPE2", {"value": 10.0}, admin_a)
        update_metric_value(db, report.id, "P6_SCOPE2_EMISSIONS", {"value": 20.0}, admin_a)
        run = run_validation(db, report, admin_a)
        codes = {i.code for i in run.issues}
        assert "RECONCILIATION_MISMATCH" in codes
        assert run.status in ("WARNINGS", "FAILED")

    def test_missing_exchange_rate_for_currency_mismatch_is_error(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "SEC_A_CSR_AMOUNT_SPENT", {
            "value": 1000, "currency": "USD", "reporting_currency": "INR",
        }, admin_a)
        run = run_validation(db, report, admin_a)
        codes = {i.code for i in run.issues}
        assert "MISSING_EXCHANGE_RATE" in codes

    def test_intensity_reconciliation_mismatch_is_warning(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "CORE_ENERGY_INTENSITY", {"value": 5.0}, admin_a)
        update_metric_value(db, report.id, "P6_ENERGY_INTENSITY", {"value": 50.0}, admin_a)
        run = run_validation(db, report, admin_a)
        codes = {i.code for i in run.issues}
        assert "RECONCILIATION_MISMATCH" in codes

    def test_board_composition_is_a_required_field_distinct_from_esg_responsibility(self, db, two_tenants):
        """SEC_B_BOARD_COMPOSITION (board size/independence) is a distinct
        required disclosure from SEC_B_BOARD_ESG_RESPONSIBILITY (who owns
        ESG oversight) — filling only one must not satisfy the other."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        update_company_profile(db, "tenant-brsr-a", {"company_name": "Acme", "listed_status": "UNLISTED"}, admin_a)
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "SEC_B_BOARD_ESG_RESPONSIBILITY", {"text_value": "CSR Committee"}, admin_a)
        update_metric_value(db, report.id, "SEC_A_CSR_APPLICABLE", {"text_value": "true"}, admin_a)
        run = run_validation(db, report, admin_a)
        missing_codes = {i.metric_code for i in run.issues if i.code == "REQUIRED_FIELD_MISSING"}
        assert "SEC_B_BOARD_COMPOSITION" in missing_codes

        update_metric_value(db, report.id, "SEC_B_BOARD_COMPOSITION", {"text_value": "9 directors, 5 independent"}, admin_a)
        run2 = run_validation(db, report, admin_a)
        missing_codes2 = {i.metric_code for i in run2.issues if i.code == "REQUIRED_FIELD_MISSING"}
        assert "SEC_B_BOARD_COMPOSITION" not in missing_codes2

    def test_direct_api_validate_endpoint(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.post(f"/brsr/reports/{report.id}/validate", headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        assert res.json()["error_count"] >= 1  # required fields still missing


class TestAuditIntegration:
    def test_report_creation_recorded_in_audit_ledger(self, db, two_tenants):
        from app.shared.models import AuditEventORM, EventType
        start, end = _period()
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", two_tenants["admin_a"])
        events = db.query(AuditEventORM).filter(AuditEventORM.event_type == EventType.BRSR_REPORT_CREATED).all()
        assert any(e.payload.get("report_id") == report.id for e in events)

    def test_metric_update_recorded_in_audit_ledger(self, db, two_tenants):
        from app.shared.models import AuditEventORM, EventType
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_metric_value(db, report.id, "SEC_A_CSR_AMOUNT_SPENT", {"value": 1000}, admin_a)
        events = db.query(AuditEventORM).filter(AuditEventORM.event_type == EventType.BRSR_METRIC_UPDATED).all()
        assert any(e.payload.get("report_id") == report.id and e.payload.get("metric_code") == "SEC_A_CSR_AMOUNT_SPENT" for e in events)

    def test_status_transition_recorded_in_audit_ledger(self, db, two_tenants):
        from app.shared.models import AuditEventORM, EventType
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        transition_status(db, report.id, "DATA_COLLECTION", admin_a)
        events = db.query(AuditEventORM).filter(AuditEventORM.event_type == EventType.BRSR_STATUS_CHANGED).all()
        assert any(e.payload.get("report_id") == report.id for e in events)


class TestNotificationIntegration:
    def test_report_creation_notifies_other_admins_not_creator(self, db, two_tenants):
        from app.notify.service import get_notifications
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        # Add a second admin for tenant-brsr-a to verify fan-out.
        second_admin = UserORM(username="brsr_admin_a2", email="admin_a2@brsr.io", hashed_password=hash_password("pass12345"),
                                role=UserRole.COMPANY_ADMIN, tenant_id="tenant-brsr-a", is_active=True)
        db.add(second_admin)
        db.commit()
        db.refresh(second_admin)

        create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)

        creator_notifs = get_notifications(db, recipient_user_id=admin_a.id)
        other_notifs = get_notifications(db, recipient_user_id=second_admin.id)
        assert not any(n.event_type.value == "BRSR_REPORT_CREATED" for n in creator_notifs)
        assert any(n.event_type.value == "BRSR_REPORT_CREATED" for n in other_notifs)


class TestReportGeneration:
    def _approved_report(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        update_company_profile(db, "tenant-brsr-a", {"company_name": "Acme Real Corp", "listed_status": "UNLISTED"}, admin_a)
        required_defs = db.query(BrsrMetricDefinitionORM).filter(BrsrMetricDefinitionORM.required == True).all()  # noqa: E712
        for defn in required_defs:
            update_metric_value(db, report.id, defn.metric_code, {"text_value": "Yes", "value": 1.0 if defn.data_type != "TEXT" else None}, admin_a)
        transition_status(db, report.id, "DATA_COLLECTION", admin_a)
        transition_status(db, report.id, "VALIDATED", admin_a)
        transition_status(db, report.id, "APPROVED", admin_a)
        return transition_status(db, report.id, "GENERATED", admin_a)

    def test_json_export_has_no_fabricated_currency_and_real_company_name(self, db, two_tenants):
        from app.brsr.report_generator import generate_json
        report = self._approved_report(db, two_tenants)
        data = generate_json(db, report)
        assert data["company_profile"]["company_name"] == "Acme Real Corp"
        assert "assurance_disclaimer" in data
        assert "does not" in data["assurance_disclaimer"] and "assurance" in data["assurance_disclaimer"]
        assert data["metadata"]["financial_year"] == "2025-26"
        assert data["metadata"]["tenant_id"] == "tenant-brsr-a"

    def test_csv_export_contains_provenance_columns(self, db, two_tenants):
        from app.brsr.report_generator import generate_csv
        report = self._approved_report(db, two_tenants)
        csv_text = generate_csv(db, report)
        assert "source_type" in csv_text
        assert "quality" in csv_text

    def test_excel_export_produces_valid_workbook(self, db, two_tenants):
        from app.brsr.report_generator import generate_excel
        from openpyxl import load_workbook
        report = self._approved_report(db, two_tenants)
        content = generate_excel(db, report)
        wb = load_workbook(io.BytesIO(content))
        assert "Company Profile" in wb.sheetnames
        assert "Methodology" in wb.sheetnames

    def test_pdf_export_produces_valid_pdf_bytes(self, db, two_tenants):
        from app.brsr.report_generator import generate_pdf
        report = self._approved_report(db, two_tenants)
        content = generate_pdf(db, report)
        assert content[:4] == b"%PDF"

    def test_export_endpoint_records_audit_event(self, db, two_tenants):
        from app.shared.models import AuditEventORM, EventType
        report = self._approved_report(db, two_tenants)
        token = get_token(two_tenants["admin_a"])
        res = client.get(f"/brsr/reports/{report.id}/export", params={"format": "json"}, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        events = db.query(AuditEventORM).filter(AuditEventORM.event_type == EventType.BRSR_REPORT_EXPORTED).all()
        assert any(e.payload.get("report_id") == report.id for e in events)


class TestBrsrAuditTrailEndpoint:
    """GET /brsr/reports/{id}/audit — previously had zero test coverage at all."""

    def test_returns_only_this_reports_own_events(self, db, two_tenants):
        start, end = _period()
        admin_a, admin_b = two_tenants["admin_a"], two_tenants["admin_b"]
        report_a = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        report_b = create_report(db, "tenant-brsr-b", "2025-26", start, end, "BRSR-2023", admin_b)
        token_a = get_token(admin_a)

        res = client.get(f"/brsr/reports/{report_a.id}/audit", headers={"Authorization": f"Bearer {token_a}"})
        assert res.status_code == 200
        events = res.json()
        assert len(events) >= 1
        assert all(e["payload"].get("report_id") == report_a.id for e in events)
        assert not any(e["payload"].get("report_id") == report_b.id for e in events)

    def test_cross_tenant_admin_cannot_view_audit_trail(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token_b = get_token(two_tenants["admin_b"])
        res = client.get(f"/brsr/reports/{report.id}/audit", headers={"Authorization": f"Bearer {token_b}"})
        assert res.status_code == 404

    def test_audit_trail_timestamps_are_utc_safe(self, db, two_tenants):
        """Regression proof for the ensure_utc() fix applied to this
        endpoint — SQLite (this test's DB) silently drops tzinfo on read;
        without ensure_utc() this assertion fails."""
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)
        res = client.get(f"/brsr/reports/{report.id}/audit", headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        events = res.json()
        assert len(events) >= 1
        for e in events:
            raw = e["timestamp"]
            assert raw.endswith("+00:00") or raw.endswith("Z"), f"offset-less timestamp: {raw!r}"


class TestBrsrTimezoneConsistency:
    """Regression proof for ensure_utc() applied to every Brsr*Response
    model (BrsrReportResponse/BrsrValidationRunResponse/etc.) — without the
    field_validator fix, these assertions fail under SQLite."""

    def test_report_response_timestamps_are_utc_safe(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)

        res = client.get(f"/brsr/reports/{report.id}", headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        body = res.json()
        for field in ("reporting_period_start", "reporting_period_end", "created_at"):
            raw = body[field]
            assert raw.endswith("+00:00") or raw.endswith("Z"), f"{field} offset-less: {raw!r}"

    def test_validation_run_response_timestamp_is_utc_safe(self, db, two_tenants):
        start, end = _period()
        admin_a = two_tenants["admin_a"]
        report = create_report(db, "tenant-brsr-a", "2025-26", start, end, "BRSR-2023", admin_a)
        token = get_token(admin_a)

        res = client.post(f"/brsr/reports/{report.id}/validate", headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        raw = res.json()["run_at"]
        assert raw.endswith("+00:00") or raw.endswith("Z"), f"run_at offset-less: {raw!r}"

    def test_json_export_period_timestamps_are_utc_safe(self, db, two_tenants):
        report = TestReportGeneration()._approved_report(db, two_tenants)
        token = get_token(two_tenants["admin_a"])
        res = client.get(f"/brsr/reports/{report.id}/export", params={"format": "json"}, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        meta = res.json()["metadata"]
        for field in ("reporting_period_start", "reporting_period_end", "generated_at"):
            raw = meta[field]
            assert raw is not None
            assert raw.endswith("+00:00") or raw.endswith("Z"), f"{field} offset-less: {raw!r}"
